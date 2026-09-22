"""
Takes the three CSVs produced by build_tracker.py (or the current state
of your working spreadsheet, exported back out to CSV in the same
column layout) plus the original Goodreads export, and:

  1. Adds "Goodreads Bookshelf", "Number of Pages", and "Year" columns,
     matched back to the Goodreads export (ISBN first, then normalized
     title+author, then normalized title alone - each only applied when
     it resolves to exactly one Goodreads record, to avoid mismatches
     from the title normalization collapsing distinct books together;
     see normalize.norm_title's docstring).
  2. Backfills a missing ISBN from the matched Goodreads record, if any.
  3. Relabels combined sources like "Goodreads & Library" as
     "Goodreads, Library" - easier to read/filter on.
  4. On Have Read specifically: a library-sourced row's Date Started is
     really a checkout date, not a reading-start date (see
     apply_checkout_dates.py). This moves it into its own
     "Library Checkout Date" column instead, so Date Started only ever
     means an actual start-reading date.
  5. On Want to Read specifically: rows with no date info at all
     (typically library-only additions - the library list doesn't carry
     an "added" date) get sorted as a block, by the author's last name
     then title, rather than left in import order.
  6. Drops the placeholder Cover column.
  7. Writes everything to a single .xlsx workbook with one tab per
     input CSV, a bold/frozen header row, and column widths auto-fit to
     content - instead of three separate CSVs/sheets.

Requires openpyxl (pip install openpyxl).

Usage:
    python enrich_and_export.py \
        --goodreads data/goodreads_export.csv \
        --want-to-read output/want_to_read.csv \
        --currently-reading output/currently_reading.csv \
        --have-read output/have_read.csv \
        --out output/Reading_Tracker.xlsx
"""
import argparse
import csv
import re

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from normalize import norm_title, norm_author
from parse_goodreads import clean_isbn

OUTPUT_COLUMNS = {
    "Want to Read": [
        "Title", "Author", "Format", "Source", "Date Added", "Date Started",
        "Date Finished", "Rating (1-5)", "ISBN", "Goodreads Bookshelf",
        "Number of Pages", "Year", "Notes",
    ],
    "Currently Reading": [
        "Title", "Author", "Format", "Source", "Date Added", "Date Started",
        "Date Finished", "Rating (1-5)", "ISBN", "Goodreads Bookshelf",
        "Number of Pages", "Year", "Notes",
    ],
    "Have Read": [
        "Title", "Author", "Format", "Source", "Date Added", "Date Started",
        "Date Finished", "Library Checkout Date", "Rating (1-5)", "ISBN",
        "Goodreads Bookshelf", "Number of Pages", "Year", "Notes",
    ],
}

# Surname prefixes that belong WITH the following word when sorting by
# last name (e.g. "Ursula K. Le Guin" -> "Le Guin", not "Guin").
SURNAME_PREFIXES = {
    "le", "la", "de", "del", "della", "van", "von", "der", "den",
    "el", "al", "mac", "mc", "di", "du", "st", "st.", "da", "dos", "das",
}


def load_input_csv(path):
    """Reads a CSV in build_tracker.py's output layout. Strips the
    leading apostrophe build_tracker.py adds to ISBNs (so spreadsheet
    apps don't mangle them as numbers)."""
    with open(path, newline="", encoding="utf-8") as f:
        rows = list(csv.DictReader(f))
    for r in rows:
        if r.get("ISBN", "").startswith("'"):
            r["ISBN"] = r["ISBN"][1:]
    return rows


def load_goodreads_lookup(path):
    """Indexes the full Goodreads export by ISBN and by normalized
    (title, author) / title alone, for enrichment matching. Returns
    dicts of {key: record} where a key present with more than one
    distinct record is dropped, so ambiguous matches are skipped rather
    than guessed at."""
    with open(path, newline="", encoding="utf-8") as f:
        rows = list(csv.DictReader(f))

    by_isbn = {}
    by_title_author = {}
    by_title = {}
    ambiguous_ta = set()
    ambiguous_t = set()

    for row in rows:
        isbn = clean_isbn(row.get("ISBN13")) or clean_isbn(row.get("ISBN"))
        year = (row.get("Original Publication Year") or "").strip() \
            or (row.get("Year Published") or "").strip()
        record = {
            "isbn": isbn,
            "bookshelf": (row.get("Bookshelves") or "").strip(),
            "pages": (row.get("Number of Pages") or "").strip(),
            "year": year,
        }
        if isbn:
            by_isbn[isbn] = record

        nt = norm_title(row.get("Title", ""))
        na = norm_author(row.get("Author", ""))
        if nt:
            ta_key = (nt, na)
            if ta_key in by_title_author:
                ambiguous_ta.add(ta_key)
            else:
                by_title_author[ta_key] = record
            if nt in by_title:
                ambiguous_t.add(nt)
            else:
                by_title[nt] = record

    for key in ambiguous_ta:
        del by_title_author[key]
    for key in ambiguous_t:
        del by_title[key]

    return by_isbn, by_title_author, by_title


def find_goodreads_record(title, author, isbn, by_isbn, by_title_author, by_title):
    if isbn and isbn in by_isbn:
        return by_isbn[isbn]
    key = (norm_title(title), norm_author(author))
    if key in by_title_author:
        return by_title_author[key]
    nt = norm_title(title)
    if nt in by_title:
        return by_title[nt]
    return None


def enrich_rows(rows, by_isbn, by_title_author, by_title, stats):
    for row in rows:
        if "Goodreads" not in row.get("Source", ""):
            continue
        isbn = (row.get("ISBN") or "").strip()
        gr = find_goodreads_record(row.get("Title", ""), row.get("Author", ""),
                                    isbn, by_isbn, by_title_author, by_title)
        if not gr:
            continue
        row["Goodreads Bookshelf"] = gr["bookshelf"]
        row["Number of Pages"] = gr["pages"]
        row["Year"] = gr["year"]
        stats["enriched"] += 1
        if not isbn and gr["isbn"]:
            row["ISBN"] = gr["isbn"]
            stats["isbn_backfilled"] += 1


def relabel_combined_sources(rows):
    for row in rows:
        row["Source"] = row.get("Source", "").replace(" & ", ", ")


def split_library_checkout_date(rows):
    """Have Read only: a library-sourced Date Started is really a
    checkout date. Move it to its own column."""
    for row in rows:
        checkout_date = ""
        if "Library" in row.get("Source", ""):
            started = (row.get("Date Started") or "").strip()
            if started:
                checkout_date = started
                row["Date Started"] = ""
        row["Library Checkout Date"] = checkout_date


def last_name_key(author):
    a = (author or "").strip()
    if not a:
        return ""
    if "," in a:
        return a.split(",")[0].strip().lower()
    tokens = a.split()
    if not tokens:
        return ""
    if len(tokens) >= 2 and tokens[-2].lower().strip(".") in SURNAME_PREFIXES:
        return " ".join(tokens[-2:]).lower()
    return tokens[-1].lower()


def title_sort_key(title):
    t = (title or "").strip().lower()
    return re.sub(r"^(the|a|an)\s+", "", t)


def sort_dateless_tail(rows):
    """Want to Read only: rows with no date info at all (typically
    library-only additions) are grouped at the tail of the import; sort
    just that block by author last name, then title, leaving dated rows
    in their existing order."""
    def has_no_dates(r):
        return not any((r.get(f) or "").strip()
                        for f in ("Date Added", "Date Started", "Date Finished"))

    dateless_idx = [i for i, r in enumerate(rows) if has_no_dates(r)]
    if not dateless_idx:
        return
    # Only touch it if the dateless rows are already contiguous at the
    # tail - if they're scattered, sorting just those indices in place
    # could reorder them relative to dated rows in a confusing way.
    if dateless_idx != list(range(dateless_idx[0], dateless_idx[0] + len(dateless_idx))):
        return
    if dateless_idx[-1] != len(rows) - 1:
        return

    start = dateless_idx[0]
    block = sorted(rows[start:], key=lambda r: (last_name_key(r.get("Author")),
                                                 title_sort_key(r.get("Title"))))
    rows[start:] = block


def build_workbook(sheets, out_path):
    wb = Workbook()
    wb.remove(wb.active)

    for tab_name, rows in sheets:
        cols = OUTPUT_COLUMNS[tab_name]
        ws = wb.create_sheet(title=tab_name)
        ws.append(cols)
        for cell in ws[1]:
            cell.font = Font(bold=True)
        for row in rows:
            ws.append([row.get(c, "") for c in cols])
        ws.freeze_panes = "A2"

        for i, col_name in enumerate(cols, start=1):
            max_len = len(str(col_name))
            for row in rows:
                val = row.get(col_name) or ""
                max_len = max(max_len, len(str(val)))
            ws.column_dimensions[get_column_letter(i)].width = min(max(max_len + 2, 8), 60)

    wb.save(out_path)


def main():
    ap = argparse.ArgumentParser(description=__doc__,
                                  formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--goodreads", required=True, help="Goodreads library export CSV")
    ap.add_argument("--want-to-read", required=True)
    ap.add_argument("--currently-reading", required=True)
    ap.add_argument("--have-read", required=True)
    ap.add_argument("--out", default="Reading_Tracker.xlsx")
    args = ap.parse_args()

    by_isbn, by_title_author, by_title = load_goodreads_lookup(args.goodreads)

    want = load_input_csv(args.want_to_read)
    current = load_input_csv(args.currently_reading)
    have = load_input_csv(args.have_read)

    stats = {"enriched": 0, "isbn_backfilled": 0}
    for rows in (want, current, have):
        enrich_rows(rows, by_isbn, by_title_author, by_title, stats)
        relabel_combined_sources(rows)

    split_library_checkout_date(have)
    sort_dateless_tail(want)

    build_workbook(
        [("Want to Read", want), ("Currently Reading", current), ("Have Read", have)],
        args.out,
    )

    print(f"Enriched {stats['enriched']} rows from the Goodreads export "
          f"({stats['isbn_backfilled']} ISBNs backfilled)")
    print(f"Wrote {args.out}: Want to Read {len(want)}, "
          f"Currently Reading {len(current)}, Have Read {len(have)}")


if __name__ == "__main__":
    main()
